from __future__ import annotations

import io
import re

SUPPORTED = {
    "application/pdf": "pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    "text/plain": "txt",
    "text/markdown": "md",
    "text/csv": "csv",
}

EXT_MAP = {
    ".pdf": "pdf",
    ".docx": "docx",
    ".txt": "txt",
    ".md": "md",
    ".markdown": "md",
    ".csv": "csv",
}


class ExtractionError(Exception):
    pass


def kind_for(filename: str, content_type: str) -> str:
    lowered = filename.lower()
    for ext, kind in EXT_MAP.items():
        if lowered.endswith(ext):
            return kind
    if content_type in SUPPORTED:
        return SUPPORTED[content_type]
    raise ExtractionError(
        "Unsupported file type. Upload a PDF, Word (.docx), Markdown, CSV, or plain text file."
    )


def _clean(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract(data: bytes, filename: str, content_type: str) -> list[tuple[int | None, str]]:
    """Returns [(page_number_or_None, text)] so citations can point at a page."""
    kind = kind_for(filename, content_type)

    if kind == "pdf":
        from pypdf import PdfReader

        try:
            reader = PdfReader(io.BytesIO(data))
        except Exception as exc:  # noqa: BLE001 - surfaced to the uploader
            raise ExtractionError(f"Could not read this PDF: {exc}") from exc
        if reader.is_encrypted:
            raise ExtractionError(
                "This PDF is password protected. Remove the password and re-upload."
            )
        pages = []
        for i, page in enumerate(reader.pages, start=1):
            text = _clean(page.extract_text() or "")
            if text:
                pages.append((i, text))
        if not pages:
            raise ExtractionError(
                "No selectable text found — this looks like a scanned PDF. "
                "Run it through OCR first, or upload a text version."
            )
        return pages

    if kind == "docx":
        import docx

        try:
            document = docx.Document(io.BytesIO(data))
        except Exception as exc:  # noqa: BLE001
            raise ExtractionError(f"Could not read this Word document: {exc}") from exc
        parts = [p.text for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        text = _clean("\n".join(parts))
        if not text:
            raise ExtractionError("This Word document appears to be empty.")
        return [(None, text)]

    try:
        text = _clean(data.decode("utf-8"))
    except UnicodeDecodeError:
        text = _clean(data.decode("latin-1", errors="replace"))
    if not text:
        raise ExtractionError("This file appears to be empty.")
    return [(None, text)]
